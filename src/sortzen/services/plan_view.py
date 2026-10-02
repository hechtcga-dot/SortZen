"""The plan as rows for the window and the Excel export: Ready, Review, Staying and folders."""
from __future__ import annotations

import os
from dataclasses import dataclass, field

from ..engine.plan import (FOLDER_REVIEW, KEEP_TOGETHER, MOVE, REVIEW, SORT_INSIDE, STAY, STAYS, Plan, Reason)
from ..scanning.file_types import kind_of

KIND_NAMES = {"word": "Word", "pdf": "PDF", "spreadsheet": "Spreadsheet", "presentation": "Presentation",
              "form": "Form", "text": "Text", "image": "Picture", "video": "Video", "audio": "Music",
              "archive": "Zip", "installer": "Program", "ebook": "E-book", "web page": "Web page",
              "shortcut": "Shortcut", "other": "Other"}
MAX_PATH = 259                  # longest full path Windows programs reliably handle

READY, TO_REVIEW, STAYING, SORTED_INSIDE = "Ready", "Review", "Staying", "Sorted from the inside"
GROUPS = (READY, TO_REVIEW, STAYING, SORTED_INSIDE)


@dataclass
class PlanRow:
    path: str
    is_folder: bool
    current: str
    destination: str | None
    percent: int
    action: str                     # "Move", "Stay", "Review", "Keep together", "Sort inside"
    reasons: list[Reason] = field(default_factory=list)
    runner_up: tuple[str, int] | None = None
    note: str = ""
    topic: str = ""
    new_folder: bool = False
    files: int = 0
    problems: list[str] = field(default_factory=list)

    @property
    def name(self) -> str:
        return os.path.basename(self.path)

    @property
    def kind(self) -> str:
        if self.is_folder:
            return "Folder"
        return KIND_NAMES.get(kind_of(os.path.splitext(self.path)[1]), "Other")

    @property
    def moves(self) -> bool:
        return self.action in ("Move", "Keep together") and bool(self.destination)


def rows(plan: Plan, autonomy: int, ask_everything: bool = False) -> dict[str, list[PlanRow]]:
    """Every file and folder of the plan in its group. Nothing in Ready when users want to be asked."""
    level = 101 if ask_everything else autonomy
    groups: dict[str, list[PlanRow]] = {g: [] for g in GROUPS}
    for f in plan.folders:
        parent = os.path.dirname(f.path)
        if f.outcome == SORT_INSIDE:
            groups[SORTED_INSIDE].append(PlanRow(f.path, True, parent, None, f.percent, "Sort inside", f.reasons,
                                                 files=f.files, topic=f.topic))
        elif f.outcome == KEEP_TOGETHER and f.destination:
            row = PlanRow(f.path, True, parent, f.destination, f.percent, "Keep together", f.reasons, files=f.files,
                          topic=f.topic, new_folder=not os.path.isdir(f.destination))
            groups[READY if f.percent >= level else TO_REVIEW].append(row)
        elif f.outcome in (FOLDER_REVIEW, KEEP_TOGETHER):
            groups[TO_REVIEW].append(PlanRow(f.path, True, parent, None, f.percent, "Review", f.reasons,
                                             files=f.files, topic=f.topic))
        elif f.outcome == STAYS and f.topic:
            groups[STAYING].append(PlanRow(f.path, True, parent, parent, f.percent, "Stay", f.reasons,
                                           files=f.files, topic=f.topic))
    for s in plan.files:
        row = PlanRow(s.path, False, s.current_folder, s.destination, s.percent,
                      {MOVE: "Move", STAY: "Stay", REVIEW: "Review"}[s.action], s.reasons, s.runner_up, s.note,
                      s.topic, s.new_folder)
        if s.action == STAY:
            groups[STAYING].append(row)
        elif s.action == MOVE and s.percent >= level:
            groups[READY].append(row)
        else:
            groups[TO_REVIEW].append(row)
    for group in (groups[READY], groups[TO_REVIEW]):
        for row in group:
            row.problems = problems(row)
    for group in groups.values():
        group.sort(key=lambda r: (-r.percent, r.path.lower()) if r.action != "Review" else (r.percent, r.path.lower()))
    return groups


def problems(row: PlanRow) -> list[str]:
    """What would stop or change this move, found before anything moves."""
    if not row.moves:
        return []
    found = []
    target = os.path.join(row.destination, row.name)
    if len(target) > MAX_PATH:
        found.append(f"The new path would be too long for Windows ({len(target)} characters)")
    if os.path.exists(target) and os.path.normcase(target) != os.path.normcase(row.path):
        stem, ext = os.path.splitext(row.name)
        found.append(f"“{row.name}” is already there: this one would be saved as “{stem} (2){ext}”")
    existing = row.destination
    while existing and not os.path.isdir(existing) and os.path.dirname(existing) != existing:
        existing = os.path.dirname(existing)
    if existing and os.path.isdir(existing) and not os.access(existing, os.W_OK):
        found.append("SortZen can't write to the destination folder")
    return found


def display(path: str | None, roots: list[str]) -> str:
    """A path as users see it: from the added folder's name down, e.g. "Downloads/older downloads"."""
    if not path:
        return ""
    for root in sorted(roots, key=len, reverse=True):
        try:
            rel = os.path.relpath(path, os.path.dirname(root))
        except ValueError:          # another drive
            continue
        if not rel.startswith(".."):
            return rel.replace(os.sep, "/")
    return path


def export_xlsx(plan: Plan, target: str, roots: list[str], autonomy: int, ask_everything: bool = False) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Plan"
    header = ["Group", "Name", "Type", "From", "To", "Action", "Sure %", "New folder", "Topic", "Reasons", "Problems"]
    sheet.append(header)
    for group, items in rows(plan, autonomy, ask_everything).items():
        for r in items:
            reasons = "\n".join(("+ " if x.supports else "- ") + x.text for x in r.reasons)
            sheet.append([group, r.name, f"Folder ({r.files} files)" if r.is_folder else r.kind,
                          display(r.current, roots), display(r.destination, roots), r.action, r.percent,
                          "Yes" if r.new_folder else "", r.topic, reasons, "\n".join(r.problems)])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJK", (14, 40, 14, 40, 40, 14, 8, 11, 18, 80, 50)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=10, max_col=10):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    questions = book.create_sheet("Questions")
    questions.append(["Question", "Choices", "Answer"])
    for q in plan.questions:
        questions.append([q.text, "\n".join(c.label for c in q.choices),
                          q.choices[q.answer].label if q.answer is not None else ""])
    for cell in questions[1]:
        cell.font = Font(bold=True)
    questions.column_dimensions["A"].width = 80
    questions.column_dimensions["B"].width = 60
    book.save(target)
