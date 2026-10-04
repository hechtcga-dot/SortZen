import os
import tempfile
import time
import unittest
from pathlib import Path

from sortzen.ai.web_providers import OllamaProvider
from sortzen.config import AppPaths
from sortzen.services import AppService
from tests.test_repositories import FakeKeyring
from sortzen.repositories.api_keys import ApiKeyStore


class AppServiceTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.paths = AppPaths(Path(self.dir.name))
        self.service = AppService(self.paths, ApiKeyStore(FakeKeyring()))

    def tearDown(self):
        self.dir.cleanup()

    def test_gemini_is_the_default(self):
        self.assertEqual(self.service.ai_service(), "gemini")

    def test_choose_service(self):
        self.service.set_ai_service("claude")
        self.assertEqual(AppService(self.paths, self.service.keys).ai_service(), "claude")
        with self.assertRaises(ValueError):
            self.service.set_ai_service("nonsense")

    def test_unknown_saved_service_falls_back_to_default(self):
        self.service.settings.set("ai_service", "gone")
        self.assertEqual(self.service.ai_service(), "gemini")

    def test_model_override_and_reset(self):
        default = self.service.model()
        self.service.set_model("custom-model")
        self.assertEqual(self.service.model(), "custom-model")
        self.service.set_model("")
        self.assertEqual(self.service.model(), default)

    def test_keys_never_written_to_settings(self):
        self.service.save_api_key("secret-key-value")
        self.service.set_model("custom-model")
        self.assertTrue(self.service.has_api_key())
        self.assertNotIn("secret-key-value", self.paths.settings_path.read_text(encoding="utf-8"))

    def test_ollama_needs_no_key(self):
        self.service.set_ai_service("ollama")
        self.assertTrue(self.service.has_api_key())
        self.assertIsInstance(self.service.provider(), OllamaProvider)

    def test_missing_key_is_a_plain_error(self):
        self.service.set_ai_service("claude")
        with self.assertRaises(ValueError):
            self.service.provider()


class FoldersTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        root = Path(self.dir.name)
        self.service = AppService(AppPaths(root / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.downloads, self.documents = root / "Downloads", root / "Documents"
        self.downloads.mkdir()
        self.documents.mkdir()

    def tearDown(self):
        self.dir.cleanup()

    def test_add_change_and_remove(self):
        from sortzen.engine.planner import TIDY
        from sortzen.services.app_service import FolderError

        self.service.add_source(str(self.downloads))
        self.service.add_destination(str(self.documents))
        with self.assertRaises(FolderError):
            self.service.add_destination(str(self.documents))
        with self.assertRaises(FolderError):
            self.service.add_source(str(self.downloads / "missing"))
        self.service.set_source_mode(str(self.downloads), TIDY)
        again = AppService(self.service.paths, self.service.keys)
        self.assertEqual(again.source_folders(), [{"path": str(self.downloads), "mode": TIDY}])
        self.assertEqual(again.destination_folders(), [str(self.documents)])
        again.remove_folder(str(self.downloads))
        self.assertEqual(again.source_folders(), [])

    def test_protected_folders_refused(self):
        from sortzen.repositories.file_index import path_key
        from sortzen.services.app_service import FolderError

        self.service.scanner.protected = [path_key(self.documents)]
        with self.assertRaises(FolderError):
            self.service.add_destination(str(self.documents))

    def test_options(self):
        self.assertFalse(self.service.option("gentle"))
        self.assertTrue(self.service.option("stop_reading_learned"))
        self.service.set_option("gentle", True)
        self.assertTrue(AppService(self.service.paths, self.service.keys).option("gentle"))
        with self.assertRaises(KeyError):
            self.service.set_option("nonsense", True)

    def test_autonomy_answers_corrections(self):
        self.assertEqual(self.service.autonomy(), 90)
        self.service.set_autonomy(120)
        self.assertEqual(self.service.autonomy(), 100)
        self.service.save_answers({"topic:x": 1})
        self.service.save_answers({"topic:y": 0, "topic:x": None})
        self.assertEqual(self.service.answers(), {"topic:y": 0})
        previous = self.service.correct(["a.pdf"], str(self.documents))
        self.assertEqual(self.service.corrections(), {"a.pdf": str(self.documents)})
        self.service.restore_corrections(previous)
        self.assertEqual(self.service.corrections(), {})

    def test_outermost_destination_only(self):
        from sortzen.services.app_service import _outermost
        a, b, c = self.documents, self.documents / "Word", Path(self.dir.name) / "Documents2"
        self.assertEqual(_outermost([b, a, c]), [a, c])


class MakePlanTest(unittest.TestCase):
    """The whole path on the test folders: add folders, plan, group, export."""

    @classmethod
    def setUpClass(cls):
        from tests.fixtures import shared_test_folders
        from sortzen.engine.planner import TIDY

        cls.dir = tempfile.TemporaryDirectory()
        root = shared_test_folders()
        cls.service = AppService(AppPaths(Path(cls.dir.name) / "data"), ApiKeyStore(FakeKeyring()))
        cls.service.scanner.protected = []
        cls.service.add_source(str(root / "Downloads"))
        cls.service.add_source(str(root / "My Drive"), TIDY)
        cls.service.add_destination(str(root / "Sorted"))
        cls.plan = cls.service.make_plan()

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def test_plan_groups(self):
        groups = self.service.plan_rows(self.plan)
        self.assertGreater(len(groups["Ready"]), 150)
        self.assertGreater(len(groups["Review"]), 20)
        self.assertTrue(any(r.is_folder and r.action == "Keep together" for r in groups["Ready"] + groups["Review"]))
        self.assertTrue(any(r.name == "older downloads" for r in groups["Sorted from the inside"]))
        self.service.set_ask_everything(True)
        try:
            self.assertEqual(self.service.plan_rows(self.plan)["Ready"], [])
        finally:
            self.service.set_ask_everything(False)

    def test_display_paths(self):
        row = self.service.plan_rows(self.plan)["Ready"][0]
        self.assertFalse(os.path.isabs(self.service.display(row.current)))

    def test_export(self):
        from openpyxl import load_workbook

        target = os.path.join(self.dir.name, "plan.xlsx")
        self.service.export_plan(self.plan, target)
        book = load_workbook(target)
        self.assertEqual(book["Plan"]["A1"].value, "Group")
        self.assertGreater(book["Plan"].max_row, 300)
        self.assertGreaterEqual(book["Questions"].max_row, 2)

    def test_destination_choices(self):
        choices = self.service.destination_choices(self.plan)
        self.assertTrue(any(c.endswith(os.path.join("Work", "Payroll")) for c in choices))

    def test_needs_a_folder_to_sort(self):
        from sortzen.services.app_service import FolderError

        empty = AppService(AppPaths(Path(self.dir.name) / "empty"), ApiKeyStore(FakeKeyring()))
        with self.assertRaises(FolderError):
            empty.make_plan()


class LeftOutServiceTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        root = Path(self.dir.name)
        self.service = AppService(AppPaths(root / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.down = root / "Downloads"
        (self.down / "Keep me").mkdir(parents=True)
        for i in range(120):
            (self.down / "Keep me" / f"note {i}.txt").write_text(f"budget note {i}")
        (self.down / "loose.txt").write_text("budget note loose")
        self.service.add_source(str(self.down))

    def tearDown(self):
        self.dir.cleanup()

    def test_left_out_round_trip_and_plan(self):
        keep = str(self.down / "Keep me")
        previous = self.service.set_left_out([keep], True)
        self.assertTrue(self.service.is_left_out(keep + os.sep + "note 1.txt"))
        plan = self.service.make_plan()
        self.assertFalse(any("Keep me" in s.path for s in plan.files))
        self.service.restore_left_out(previous)
        self.assertFalse(self.service.is_left_out(keep))

    def test_learned_folders_are_read_by_name_only(self):
        from unittest import mock
        from sortzen.scanning import scanner as scanner_module

        keep = str(self.down / "Keep me")
        self.service.set_left_out([keep], True)
        self.service.make_plan()                                   # first run reads everything
        self.assertEqual(self.service.learned_folders(self.service.count_folders()), [keep])
        (self.down / "Keep me" / "new.txt").write_text("budget note new")
        with mock.patch.object(scanner_module, "read_contents", side_effect=AssertionError("opened")):
            self.service.make_plan()
        self.service.set_option("stop_reading_learned", False)
        self.assertEqual(self.service.learned_folders(self.service.count_folders()), [])

    def test_count_tree_and_estimate(self):
        counts = self.service.count_folders()
        self.assertEqual(counts[0].files, 121)
        self.assertEqual(counts[0].tree[str(self.down / "Keep me")][0], 120)
        self.assertGreater(self.service.estimate_seconds(counts), 0)


class ProblemsTest(unittest.TestCase):
    def test_problems_found_before_moving(self):
        from sortzen.services.plan_view import PlanRow, problems

        with tempfile.TemporaryDirectory() as d:
            dest = os.path.join(d, "Work")
            os.mkdir(dest)
            Path(dest, "report.pdf").write_text("x")
            clash = PlanRow(os.path.join(d, "Downloads", "report.pdf"), False, os.path.join(d, "Downloads"),
                            dest, 95, "Move")
            self.assertIn("would be saved as “report (2).pdf”", " ".join(problems(clash)))
            deep = os.path.join(d, *["a long folder name"] * 15)
            long_row = PlanRow(os.path.join(d, "x.pdf"), False, d, deep, 95, "Move")
            self.assertIn("too long for Windows", " ".join(problems(long_row)))
            staying = PlanRow(os.path.join(d, "x.pdf"), False, d, d, 95, "Stay")
            self.assertEqual(problems(staying), [])


class MoveServiceTest(unittest.TestCase):
    """Moving ticked rows of a real plan on a copy of the test folders, then putting them back."""

    def setUp(self):
        import shutil

        from tests.fixtures import shared_test_folders

        self.dir = tempfile.TemporaryDirectory()
        self.root = Path(self.dir.name) / "folders"
        for name in ("Downloads", "Sorted"):
            shutil.copytree(shared_test_folders() / name, self.root / name)
        self.service = AppService(AppPaths(Path(self.dir.name) / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(self.root / "Downloads"))
        self.service.add_destination(str(self.root / "Sorted"))
        self.plan = self.service.make_plan()

    def tearDown(self):
        self.dir.cleanup()

    def snapshot(self):
        return sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*"))

    def test_preview_move_remember_and_undo(self):
        ready = self.service.plan_rows(self.plan)["Ready"]
        before = self.snapshot()
        preview = self.service.move_preview(self.plan, ready)
        self.assertEqual(preview.items, len(ready))
        self.assertEqual(sum(n for _, n, _ in preview.destinations), preview.items)
        self.assertEqual(self.snapshot(), before)                  # the preview touches nothing
        result = self.service.move(self.plan, ready)
        self.assertEqual((result.moved, result.failed), (len(ready), []))
        for row in ready:
            self.assertFalse(os.path.exists(row.path))
            self.assertTrue(os.path.isdir(row.destination))
        self.assertEqual(self.service.move_runs()[0]["moved"], len(ready))
        remembered = {r.path for r in self.service.index.records(self.root / "Sorted")}
        self.assertTrue(all(target in remembered for _, target in result.moves if os.path.isfile(target)))
        again = self.service.make_plan()                           # moved files are remembered, not read again
        self.assertNotIn(ready[0].path, {s.path for s in again.files})
        undone = self.service.undo_move(result.log)
        self.assertEqual((undone.moved, undone.failed), (len(ready), []))
        self.assertEqual(self.snapshot(), before)

    def test_files_inside_a_moving_folder_go_with_it(self):
        from sortzen.services import moving
        from sortzen.services.plan_view import PlanRow

        folder = self.root / "Downloads" / "Kestrel bid"
        (folder / "a.txt").parent.mkdir(parents=True, exist_ok=True)
        (folder / "a.txt").write_text("x")
        rows = [PlanRow(str(folder), True, str(folder.parent), str(self.root / "Sorted"), 95, "Keep together"),
                PlanRow(str(folder / "a.txt"), False, str(folder), str(self.root / "Sorted"), 95, "Move")]
        self.assertEqual([r.path for r in moving.requests(rows)], [str(folder)])


class CopiesTest(unittest.TestCase):
    """Exact copies: found, the right one kept, extras queued in dated folders, Undo."""

    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.root = Path(self.dir.name) / "folders"
        self.downloads = self.root / "Downloads"
        self.sorted = self.root / "Sorted"
        for path, text, when in (("Sorted/Taxes/2024 T4.pdf", "slip", 100), ("Downloads/2024 T4.pdf", "slip", 50),
                                 ("Downloads/old/2024 T4 (1).pdf", "slip", 200), ("Downloads/a.txt", "only one", 1),
                                 ("Downloads/b (1).txt", "twin", 300), ("Downloads/b.txt", "twin", 400)):
            target = self.root / path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(text, encoding="utf-8")
            os.utime(target, (1_700_000_000 + when, 1_700_000_000 + when))
        self.service = AppService(AppPaths(Path(self.dir.name) / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(self.downloads))
        self.service.add_destination(str(self.sorted))
        self.plan = self.service.make_plan()

    def tearDown(self):
        self.dir.cleanup()

    def group(self, name):
        return next(g for g in self.plan.copies if any(c.name == name for c in g.copies))

    def test_copies_found_and_the_right_one_kept(self):
        self.assertEqual(len(self.plan.copies), 2)
        slip = self.group("2024 T4.pdf")
        self.assertEqual(slip.kept.path, str(self.sorted / "Taxes" / "2024 T4.pdf"))
        self.assertIn("It's in an organised folder", slip.kept.reasons)
        self.assertTrue(all(c.ticked for c in slip.extras))
        twin = self.group("b.txt")
        self.assertEqual(twin.kept.name, "b.txt")                     # no copy number beats older
        self.assertEqual(twin.kept.reasons, ["Its name has no copy number"])

    def test_queue_check_and_undo(self):
        slip = self.group("2024 T4.pdf")
        (self.downloads / "old" / "2024 T4 (1).pdf").write_text("changed", encoding="utf-8")
        result = self.service.queue_copies(self.plan.copies)
        queue = self.downloads / "To delete"
        self.assertEqual(result.moved, 2)
        self.assertTrue((queue / "2024 T4.pdf").exists())
        self.assertTrue((queue / "b (1).txt").exists())
        self.assertIn("no longer an exact copy", result.failed[0][1])     # changed since: stays
        self.assertTrue((self.downloads / "old" / "2024 T4 (1).pdf").exists())
        self.assertTrue(slip.kept.path and os.path.exists(slip.kept.path))
        again = self.service.make_plan()                                  # queued copies are never read or sorted
        self.assertFalse(any("To delete" in s.path for s in again.files))
        self.assertEqual(self.service.move_runs()[0]["kind"], "duplicates")
        self.service.undo_move(result.log)
        self.assertTrue((self.downloads / "2024 T4.pdf").exists())
        self.assertFalse(any(queue.rglob("*.*")))


class ProfileTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        base = Path(self.dir.name)
        for name in ("Downloads", "Sorted/Work", "Elsewhere/Sorted/Work"):
            (base / name).mkdir(parents=True)
        self.base = base
        self.service = AppService(AppPaths(base / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(base / "Downloads"))
        self.service.add_destination(str(base / "Sorted"))
        self.service.correct([str(base / "Downloads" / "a.pdf")], str(base / "Sorted" / "Work"))
        self.service.set_autonomy(80)
        self.service.save_api_key("secret-key", "gemini")
        self.service.settings.set("speeds", {"read": 1.0, "remembered": 2.0, "plan": 3.0})

    def tearDown(self):
        self.dir.cleanup()

    def other_pc(self):
        return AppService(AppPaths(self.base / "data2"), ApiKeyStore(FakeKeyring()))

    def test_round_trip_without_keys_and_with_a_moved_folder(self):
        target = str(self.base / "me.szprofile")
        self.service.save_profile(target)
        text = Path(target).read_text(encoding="utf-8")
        self.assertNotIn("secret-key", text)
        self.assertNotIn("speeds", text)
        other = self.other_pc()
        loaded = other.read_profile(target)
        shutil_target = self.base / "Elsewhere" / "Sorted"
        before = other.load_profile(loaded, {str(self.base / "Sorted"): str(shutil_target)})
        self.assertEqual(other.destination_folders(), [str(shutil_target)])
        self.assertEqual(other.corrections(), {str(self.base / "Downloads" / "a.pdf"): str(shutil_target / "Work")})
        self.assertEqual(other.autonomy(), 80)
        self.assertFalse(other.has_api_key("gemini"))
        other.restore_settings(before)
        self.assertEqual(other.destination_folders(), [])

    def test_keys_only_when_asked_missing_and_dropped_folders(self):
        target = str(self.base / "with keys.szprofile")
        self.service.save_profile(target, include_keys=True)
        other = self.other_pc()
        loaded = other.read_profile(target)
        (self.base / "Sorted" / "Work").rmdir()
        (self.base / "Sorted").rmdir()
        self.assertEqual(other.missing_folders(loaded), [str(self.base / "Sorted")])
        other.load_profile(loaded, dropped=[str(self.base / "Sorted")])
        self.assertEqual(other.destination_folders(), [])
        self.assertEqual(other.api_key("gemini"), "secret-key")

    def test_not_a_profile(self):
        bad = self.base / "notes.szprofile"
        bad.write_text('{"hello": 1}', encoding="utf-8")
        with self.assertRaises(ValueError):
            self.service.read_profile(str(bad))

    def test_diagnostics_have_no_paths(self):
        self.service.log_path().parent.mkdir(parents=True, exist_ok=True)
        self.service.log_path().write_text(f"2026-10-02 ERROR move stopped: can't read {self.base}/Downloads/x.pdf\n",
                                           encoding="utf-8")
        info = self.service.diagnostics()
        self.assertIn("SortZen", info)
        self.assertIn("1 to sort", info)
        self.assertNotIn(str(self.base), info)
        self.assertNotIn("secret-key", info)


class RulesRecentsRenameTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        base = Path(self.dir.name)
        self.downloads, self.sorted = base / "Downloads", base / "Sorted"
        files = {"Downloads/Northgate invoice 1.pdf": "invoice from northgate", "Downloads/Prairie invoice 2.pdf":
                 "invoice from prairie", "Downloads/Clearwater invoice 3.pdf": "invoice from clearwater",
                 "Downloads/holiday notes.txt": "beach trip", "Sorted/Trips/beach plan.txt": "beach trip plan",
                 "Sorted/Invoices/readme.txt": "kept here", "Sorted/Other/misc.txt": "misc"}
        for rel, text in files.items():
            (base / rel).parent.mkdir(parents=True, exist_ok=True)
            (base / rel).write_text(text, encoding="utf-8")
        self.service = AppService(AppPaths(base / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(self.downloads))
        self.service.add_destination(str(self.sorted))
        self.plan = self.service.make_plan()

    def tearDown(self):
        self.dir.cleanup()

    def test_recent_destinations_newest_first_and_capped(self):
        for i in range(25):
            folder = self.sorted / f"Box {i}"
            folder.mkdir()
            self.service.note_destination(str(folder))
        self.service.note_destination(str(self.sorted / "Box 3"))
        recent = self.service.recent_destinations()
        self.assertEqual(len(recent), 20)
        self.assertEqual(os.path.basename(recent[0]), "Box 3")
        (self.sorted / "Box 24").rmdir()
        self.assertNotIn(str(self.sorted / "Box 24"), self.service.recent_destinations())

    def test_rule_suggested_made_applied_and_undone(self):
        invoices = str(self.sorted / "Invoices")
        self.service.correct([str(self.downloads / "Northgate invoice 1.pdf"),
                              str(self.downloads / "Prairie invoice 2.pdf")], invoices)
        suggestion = self.service.suggest_rule(self.plan, invoices)      # the plan on screen when choosing
        self.assertEqual(suggestion.rule.word, "invoice")
        self.assertEqual(suggestion.matches, [str(self.downloads / "Clearwater invoice 3.pdf")])
        before = self.service.add_rule(suggestion.rule)
        plan = self.service.make_plan()
        third = plan.for_path(str(self.downloads / "Clearwater invoice 3.pdf"))
        self.assertEqual((third.destination, third.percent), (invoices, 100))
        self.assertIn("Your rule", third.reasons[0].text)
        self.assertIsNone(self.service.suggest_rule(plan, invoices))        # already a rule
        self.service.restore_rules(before)
        self.assertEqual(self.service.rules(), [])
        self.service.decline_rule(suggestion.rule)
        declined_pdf = suggestion.rule.__class__(suggestion.rule.word, invoices, ".pdf")
        self.service.decline_rule(declined_pdf)
        self.assertIsNone(self.service.suggest_rule(self.plan, invoices))

    def test_rename_a_planned_and_an_existing_folder(self):
        planned = str(self.sorted / "Receipts")
        self.service.correct([str(self.downloads / "Northgate invoice 1.pdf")], planned)
        self.service.note_destination(planned)
        self.assertIsNone(self.service.rename_folder(self.plan, planned, "Bills 2024"))
        self.assertFalse(os.path.exists(self.sorted / "Bills 2024"))       # nothing made on disk
        plan = self.service.make_plan()
        self.assertEqual(plan.for_path(str(self.downloads / "Northgate invoice 1.pdf")).destination,
                         str(self.sorted / "Bills 2024"))
        with self.assertRaises(ValueError):
            self.service.rename_folder(plan, str(self.sorted / "Other"), "Trips")       # name taken
        with self.assertRaises(ValueError):
            self.service.rename_folder(plan, str(self.sorted / "Other"), "a/b")
        self.service.correct([str(self.downloads / "holiday notes.txt")], str(self.sorted / "Other"))
        run = self.service.rename_folder(plan, str(self.sorted / "Other"), "Odds and ends")
        self.assertTrue((self.sorted / "Odds and ends" / "misc.txt").exists())
        self.assertEqual(self.service.corrections()[str(self.downloads / "holiday notes.txt")],
                         str(self.sorted / "Odds and ends"))
        self.service.undo_move(run.log)
        self.assertTrue((self.sorted / "Other" / "misc.txt").exists())
        self.assertEqual(self.service.corrections()[str(self.downloads / "holiday notes.txt")],
                         str(self.sorted / "Other"))


class GroupsServiceTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        base = Path(self.dir.name)
        self.downloads, self.sorted = base / "Downloads", base / "Sorted"
        for n in (48213, 99120, 1203, 77):
            (self.downloads).mkdir(exist_ok=True)
            (self.downloads / f"{n}.pdf").write_bytes(b"%PDF-1.4\n%" + str(n).encode())
        (self.sorted / "Recipes").mkdir(parents=True)
        (self.sorted / "Recipes" / "Lemon tart.txt").write_text("lemon tart recipe", encoding="utf-8")
        self.service = AppService(AppPaths(base / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(self.downloads))
        self.service.add_destination(str(self.sorted))
        self.plan = self.service.make_plan()

    def tearDown(self):
        self.dir.cleanup()

    def test_place_a_group_with_a_rule_then_undo(self):
        groups = self.service.file_groups(self.plan)
        self.assertEqual(groups[0].title, "4 PDFs whose names are only numbers")
        scans = str(self.sorted / "Scans")
        before = self.service.place_group(groups[0], scans, make_rule=True)
        self.assertEqual(len(self.service.corrections()), 4)
        (self.downloads / "555.pdf").write_bytes(b"%PDF-1.4\n%555")
        plan = self.service.make_plan()
        new = plan.for_path(str(self.downloads / "555.pdf"))
        self.assertEqual((new.destination, new.percent), (scans, 100))
        self.assertIn("Names that are only numbers (PDF files)", new.reasons[0].text)
        self.assertEqual(self.service.file_groups(plan), [])
        self.service.undo_place_group(before)
        self.assertEqual((self.service.corrections(), self.service.rules()), ({}, []))

    def test_typed_answer_saved_as_a_folder(self):
        self.service.save_answers({"topic:x": str(self.sorted / "Typed"), "folder:y": 1})
        self.assertEqual(self.service.answers(), {"topic:x": str(self.sorted / "Typed"), "folder:y": 1})
        self.service.save_answers({"topic:x": None})
        self.assertEqual(self.service.answers(), {"folder:y": 1})


class CatalogTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        base = Path(self.dir.name)
        self.sorted, self.downloads = base / "Sorted", base / "Downloads"
        files = {"Sorted/Work/Payroll/March pay register.txt": "pay register march payroll",
                 "Sorted/Work/Payroll/April pay register.txt": "pay register april payroll",
                 "Sorted/Work/Budget/Budget 2024.txt": "budget forecast spending plan",
                 "Sorted/Work/Budget/Budget 2025.txt": "budget forecast spending plan",
                 "Sorted/Recipes/Lemon tart.txt": "lemon tart recipe sugar",
                 "Sorted/Recipes/Apple pie.txt": "apple pie recipe sugar",
                 "Sorted/Apps/TideLog/requirements.txt": "pyside6",
                 "Sorted/Apps/TideLog/main.py": "print('hi')",
                 "Downloads/Plum cake recipe.txt": "plum cake recipe sugar",
                 "Downloads/Budget 2026.txt": "budget forecast spending plan"}
        for rel, text in files.items():
            (base / rel).parent.mkdir(parents=True, exist_ok=True)
            (base / rel).write_text(text, encoding="utf-8")
        self.service = AppService(AppPaths(base / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(self.downloads))
        self.service.add_destination(str(self.sorted))
        self.plan = self.service.make_plan()

    def tearDown(self):
        self.dir.cleanup()

    def category(self, name):
        return next(c for c in self.service.catalog() if os.path.basename(c.path) == name)

    def test_catalog_from_folders_without_programs(self):
        names = [os.path.basename(c.path) for c in self.service.catalog()]
        self.assertEqual(names, ["Sorted", "Apps", "Recipes", "Work", "Budget", "Payroll"])
        self.assertEqual(self.category("Work").files, 4)
        self.assertEqual(self.category("Apps").files, 2)                 # the program counts, but isn't a category
        self.assertEqual(self.category("Payroll").parent, str(self.sorted / "Work"))
        self.assertEqual(self.category("Recipes").examples, ["Apple pie.txt", "Lemon tart.txt"])

    def test_edits_rename_feedback_subcategory_and_folders(self):
        payroll = str(self.sorted / "Work" / "Payroll")
        self.service.rename_category(payroll, "Pay and timesheets")
        self.service.category_feedback(payroll, "too_broad", "stubs and timesheets are mixed")
        self.service.add_category_folder(payroll, str(self.downloads))
        made = self.service.add_subcategory(payroll, "Timesheets", "weekly hours")
        c = self.category("Payroll")
        self.assertEqual(c.name, "Pay and timesheets")
        self.assertEqual((c.last_feedback, c.feedback[-1]["text"]), ("too_broad", "stubs and timesheets are mixed"))
        self.assertEqual(c.folders, [payroll, str(self.downloads)])
        self.assertEqual(self.category("Timesheets").note, "weekly hours")
        with self.assertRaises(ValueError):
            self.service.add_subcategory(payroll, "Timesheets")
        self.service.remove_empty_subcategory(made)
        self.assertFalse(os.path.exists(made))

    def test_hidden_and_merged_categories_in_plans(self):
        cake = str(self.downloads / "Plum cake recipe.txt")
        budget = str(self.downloads / "Budget 2026.txt")
        self.assertEqual(self.plan.for_path(cake).destination, str(self.sorted / "Recipes"))
        self.assertEqual(self.plan.for_path(budget).destination, str(self.sorted / "Work" / "Budget"))
        self.service.set_category_hidden(str(self.sorted / "Recipes"), True)
        self.service.merge_category(str(self.sorted / "Work" / "Budget"), str(self.sorted / "Work" / "Payroll"))
        plan = self.service.make_plan()
        self.assertNotEqual(plan.for_path(cake).destination, str(self.sorted / "Recipes"))
        self.assertEqual(plan.for_path(budget).destination, str(self.sorted / "Work" / "Payroll"))
        self.assertTrue(self.category("Recipes").hidden)
        self.assertEqual(self.category("Budget").merged_into, str(self.sorted / "Work" / "Payroll"))

    def test_moving_a_categorys_files_with_undo(self):
        budget, payroll = str(self.sorted / "Work" / "Budget"), str(self.sorted / "Work" / "Payroll")
        requests = self.service.merge_files(budget, payroll)
        self.assertEqual(len(requests), 2)
        result = self.service.reorganize(requests)
        self.assertEqual(len(os.listdir(payroll)), 4)
        self.assertEqual(self.service.move_runs()[0]["kind"], "catalog")
        self.service.undo_move(result.log)
        self.assertEqual(len(os.listdir(budget)), 2)

    def test_files_put_into_a_category_are_learned_and_undone(self):
        recipes, payroll = str(self.sorted / "Recipes"), str(self.sorted / "Work" / "Payroll")
        files = self.service.category_files(recipes)
        self.assertEqual([f["name"] for f in files], ["Apple pie.txt", "Lemon tart.txt"])
        self.assertGreater(files[0]["size"], 0)
        cake = str(self.downloads / "Plum cake recipe.txt")
        work = str(self.sorted / "Work")
        requests = self.service.drop_requests([cake, files[0]["path"], payroll, work], payroll)
        self.assertEqual([r.path for r in requests], [cake, files[0]["path"]])   # not into itself or where it is
        result = self.service.place(requests)
        moved_cake = os.path.join(payroll, "Plum cake recipe.txt")
        self.assertTrue(os.path.isfile(moved_cake))
        self.assertEqual(self.service.corrections()[moved_cake], payroll)       # remembered as users' choice
        self.assertEqual(self.service.recent_destinations()[0], payroll)
        self.assertEqual(self.category("Payroll").chosen, 2)
        self.assertEqual(self.service.move_runs()[0]["kind"], "placed")
        plan = self.service.make_plan()
        self.assertEqual(plan.for_path(str(self.downloads / "Budget 2026.txt")).destination,
                         str(self.sorted / "Work" / "Budget"))
        self.service.undo_move(result.log)
        self.assertTrue(os.path.isfile(cake))
        self.assertEqual(self.service.corrections(), {})                         # the choices are forgotten
        self.assertEqual(self.category("Payroll").chosen, 0)

    def test_a_category_put_into_another_keeps_its_note_and_name(self):
        recipes, work = str(self.sorted / "Recipes"), str(self.sorted / "Work")
        self.service.set_folder_note(recipes, "baking")
        self.service.rename_category(recipes, "Baking")
        result = self.service.place(self.service.drop_requests([recipes], work))
        moved = os.path.join(work, "Recipes")
        self.assertEqual((self.category("Recipes").parent, self.category("Recipes").name), (work, "Baking"))
        self.assertEqual(self.service.folder_note(moved), "baking")
        self.service.undo_move(result.log)
        self.assertEqual(self.category("Recipes").parent, str(self.sorted))
        self.assertEqual(self.service.folder_note(recipes), "baking")

    def test_deleting_queues_files_and_undo_puts_them_back(self):
        tart = str(self.sorted / "Recipes" / "Lemon tart.txt")
        requests = self.service.delete_requests([tart, str(self.sorted / "Recipes")])
        self.assertEqual(len(requests), 1)
        self.assertEqual(requests[0].destination, str(self.sorted / "To delete" / "Recipes"))   # its place kept
        result = self.service.delete(requests)
        self.assertFalse(os.path.exists(tart))
        self.assertEqual(self.category("Recipes").files, 1)                     # queued files leave the catalog
        self.assertEqual(self.service.move_runs()[0]["kind"], "deleted")
        self.service.undo_move(result.log)
        self.assertTrue(os.path.isfile(tart))


class LabelsTest(CatalogTest):
    def test_label_list_priority_rename_and_remove(self):
        for name in ("Work", "Baking", "Taxes"):
            self.service.add_label(name)
        with self.assertRaises(ValueError):
            self.service.add_label("taxes")
        self.service.set_label_order(["Taxes", "Work"])
        self.assertEqual(self.service.labels(), ["Taxes", "Work", "Baking"])   # the first counts most
        cake = str(self.downloads / "Plum cake recipe.txt")
        self.service.set_labels([cake], "Baking", True)
        self.service.rename_label("Baking", "Cakes")
        self.assertEqual(self.service.labels_of(cake), ["Cakes"])
        self.service.remove_label("Cakes")
        self.assertEqual((self.service.labels(), self.service.labels_of(cake)), (["Taxes", "Work"], []))

    def test_guessed_labels_users_changes_and_learning(self):
        cake, budget = str(self.downloads / "Plum cake recipe.txt"), str(self.downloads / "Budget 2026.txt")
        self.service.add_label("Recipe")
        self.service.add_label("Budget")
        plan = self.service.make_plan()
        self.assertEqual(self.service.label_guesses(cake)[0][:2], ("Recipe", 85))   # its name says so
        self.assertEqual(self.service.labels_of(str(self.sorted / "Recipes" / "Apple pie.txt")), ["Recipe"])
        self.assertEqual(self.service.learned_since_plan(), 0)
        self.service.set_labels([budget], "Recipe", True)                  # guessed Budget is kept, Recipe added
        self.assertEqual(self.service.label_guesses(budget), [("Budget", 100, "You"), ("Recipe", 100, "You")])
        self.service.set_labels([budget], "Recipe", False)
        self.service.set_labels([budget], "Budget", False)
        self.assertEqual(self.service.labels_of(budget), [])                # users took every label away
        self.assertEqual(self.service.learned_since_plan(), 3)
        self.assertFalse(self.service.ready_to_recatalog())
        self.assertIn(budget, [s.path for s in self.service.unsure_files(plan)] + [budget])

    def test_labels_lead_to_folders_and_label_rules(self):
        from sortzen.engine.rules import Rule

        cake = str(self.downloads / "Plum cake recipe.txt")
        recipes = str(self.sorted / "Recipes")
        self.service.add_label("Sweet")
        for name in ("Apple pie.txt", "Lemon tart.txt"):
            self.service.set_labels([str(self.sorted / "Recipes" / name)], "Sweet", True)
        self.service.set_labels([cake], "Sweet", True)
        plan = self.service.make_plan()
        self.assertEqual(plan.for_path(cake).destination, recipes)
        self.service.add_rule(Rule("", str(self.sorted / "Work"), label="Sweet"))
        self.assertEqual(self.service.rules()[0].label, "Sweet")
        plan = self.service.make_plan()
        s = plan.for_path(cake)
        self.assertEqual((s.destination, s.percent), (str(self.sorted / "Work"), 100))
        self.assertIn("Files labelled “Sweet”", s.reasons[0].text)

    def test_file_notes_and_kept_together_folders(self):
        budget = str(self.downloads / "Budget 2026.txt")
        self.service.set_file_note(budget, "this is a recipes thing, not work")
        self.assertEqual(self.service.file_note(budget), "this is a recipes thing, not work")
        plan = self.service.make_plan()
        s = plan.for_path(budget)
        self.assertEqual(s.destination, str(self.sorted / "Recipes"))
        self.assertTrue(s.reasons[0].text.startswith("Your note:"))
        trip = self.downloads / "Trip photos"
        trip.mkdir()
        (trip / "beach.txt").write_text("beach day", encoding="utf-8")
        self.service.keep_folder_together(str(trip))
        plan = self.service.make_plan()
        folder = next(f for f in plan.folders if f.path == str(trip))
        self.assertEqual(folder.reasons[0].text, "Your answer")
        with self.assertRaises(ValueError):
            self.service.keep_folder_together(str(self.sorted / "Recipes"))

    def test_agreement_with_users_checks(self):
        cake, budget = str(self.downloads / "Plum cake recipe.txt"), str(self.downloads / "Budget 2026.txt")
        self.service.add_label("Recipe")
        self.service.add_label("Budget")
        plan = self.service.make_plan()
        self.service.confirm_labels([cake])                                 # agrees with the guess
        self.assertEqual(self.service.labels_of(cake), ["Recipe"])
        self.service.set_labels([budget], "Recipe", True)                  # SortZen hadn't guessed this
        self.assertEqual(self.service.agreement(), (1, 2))
        cake_plan = plan.for_path(cake)
        self.service.correct([cake], cake_plan.destination, {cake: cake_plan.destination})
        self.assertEqual(self.service.agreement(), (2, 3))
        self.assertFalse(self.service.learned_enough())
        self.service._check([True] * 40)
        self.assertTrue(self.service.learned_enough())

    def test_folders_suggested_for_labels(self):
        self.service.add_label("Sweet")
        self.service.add_label("Garden")
        for name in ("Apple pie.txt", "Lemon tart.txt"):
            self.service.set_labels([str(self.sorted / "Recipes" / name)], "Sweet", True)
        (self.sorted / "Recipes" / "Fudge.txt").write_text("fudge sugar sweet", encoding="utf-8")
        for n in range(5):
            (self.downloads / f"Garden bed plan {n}.txt").write_text(f"soil seeds {n}", encoding="utf-8")
        self.service.set_labels([str(self.downloads / "Plum cake recipe.txt")], "Sweet", True)
        plan = self.service.make_plan()
        self.service.set_labels([str(self.sorted / "Recipes" / "Fudge.txt")], "Sweet", True)
        plan = self.service.make_plan()
        found = {x.rule.label: x for x in self.service.label_folder_suggestions(plan)}
        self.assertIn("Garden", found)
        self.assertTrue(found["Garden"].new_folder)
        self.assertEqual(found["Garden"].rule.destination, os.path.join(str(self.sorted), "Garden"))
        self.service.add_rule(found["Garden"].rule)
        plan = self.service.make_plan()
        s = plan.for_path(str(self.downloads / "Garden bed plan 0.txt"))
        self.assertEqual((s.destination, s.percent), (os.path.join(str(self.sorted), "Garden"), 100))

    def test_ai_suggests_labels_and_labels_a_sample(self):
        import json

        from sortzen.ai.provider import AIProvider, AIResponse, TokenUsage

        self.service.make_plan()

        class Labeller(AIProvider):
            sent = []

            def generate_json(self, model, contents):
                Labeller.sent.append(contents[0])
                if "Suggest labels" in contents[0]:
                    return AIResponse(json.dumps({"labels": [{"name": "Baking", "why": "recipes"},
                                                             {"name": "Money"}, {"name": "baking"}]}),
                                      TokenUsage(300, 50, 350))
                files = [line for line in contents[0].splitlines() if line[:1].isdigit()]
                return AIResponse(json.dumps({"files": [
                    {"n": i, "labels": [{"label": "Baking" if "recipe" in line.lower() or "pie" in line.lower()
                                         or "tart" in line.lower() else "Money", "sure": 90},
                                        {"label": "Nonsense", "sure": 99}]}
                    for i, line in enumerate(files, start=1)]}), TokenUsage(500, 100, 600))

        labels = self.service.ai_suggest_labels(provider=Labeller())
        self.assertEqual(labels, [("Baking", "recipes"), ("Money", "")])
        for name, _ in labels:
            self.service.add_label(name)
        estimate = self.service.ai_label_estimate("sample")
        self.assertGreater(estimate["files"], 0)
        n = self.service.ai_label_files("sample", provider=Labeller())
        self.assertEqual(n, estimate["files"])
        cake = str(self.downloads / "Plum cake recipe.txt")
        self.assertEqual(self.service.label_guesses(cake), [("Baking", 90, "AI")])
        self.assertNotIn("plum cake recipe sugar", " ".join(Labeller.sent))    # names only by default

    def test_similar_and_different_change_percentages_only(self):
        cake, budget = str(self.downloads / "Plum cake recipe.txt"), str(self.downloads / "Budget 2026.txt")
        tart = str(self.sorted / "Recipes" / "Lemon tart.txt")
        before = self.plan.for_path(budget)
        self.service.pair_files([budget], tart, "similar", self.plan)
        after = self.plan.for_path(budget)
        self.assertIn("You said it's like “Lemon tart.txt”", " ".join(r.text for r in after.reasons))
        self.assertEqual(self.service.pairs_of(budget), [("similar", tart)])
        self.service.pair_files([cake], tart, "different")
        self.assertEqual(self.service.pairs_of(budget), [("similar", tart)])
        plan = self.service.make_plan()
        self.assertLessEqual(plan.for_path(cake).percent, 60)
        self.assertNotIn(os.path.normcase(cake), {os.path.normcase(p) for p in self.service.corrections()})
        self.service.pair_files([budget], tart, "different")                 # replaces what was said before
        self.assertEqual(self.service.pairs_of(budget), [("different", tart)])
        self.service.forget_pairs([budget, cake])
        self.assertEqual(self.service.pairs(), {"similar": [], "different": []})
        self.assertTrue(os.path.isfile(cake) and os.path.isfile(budget))   # nothing moved
        self.assertIsNotNone(before)


class CatalogSuggestionsTest(CatalogTest):
    def test_local_and_ai_suggestions_accept_and_turn_down(self):
        import json

        from sortzen.ai.provider import AIProvider, AIResponse, TokenUsage

        payroll = str(self.sorted / "Work" / "Payroll")
        for i in range(5):
            (Path(payroll) / f"Timesheet week {i}.txt").write_text("hours", encoding="utf-8")
        (Path(payroll) / "Holiday schedule.txt").write_text("days off", encoding="utf-8")
        self.service.category_feedback(payroll, "too_broad", "timesheets are mixed in")
        local = self.service.catalog_suggestions()
        split = next(s for s in local if s.kind == "split" and s.category == payroll)
        self.assertEqual(split.parts[0][0], "Timesheet")
        requests = self.service.split_requests(split)
        self.assertEqual(len(requests), 5)
        result = self.service.reorganize(requests)
        self.assertEqual(len(os.listdir(Path(payroll) / "Timesheet")), 5)
        self.service.undo_move(result.log)
        self.assertFalse((Path(payroll) / "Timesheet").exists())

        numbers = {os.path.basename(c.path): n for n, c in enumerate(self.service.catalog(), start=1)}

        class Reviewer(AIProvider):
            sent = []

            def generate_json(self, model, contents):
                Reviewer.sent.append(contents[0])
                return AIResponse(json.dumps({"suggestions": [
                    {"kind": "rename", "category": numbers["Recipes"], "name": "Cooking", "why": "clearer"},
                    {"kind": "merge", "category": numbers["Budget"], "into": numbers["Payroll"], "why": "same"},
                    {"kind": "new", "category": numbers["Work"], "name": "Contracts", "note": "signed contracts"},
                    {"kind": "rename", "category": 999, "name": "Nowhere"}]}), TokenUsage(800, 200, 1000))

        found = self.service.ask_ai_about_catalog(provider=Reviewer())
        self.assertEqual([s.kind for s in found], ["rename", "merge", "new"])
        self.assertIn("feedback: too broad - timesheets are mixed in", Reviewer.sent[0])
        self.assertNotIn("pay register march", Reviewer.sent[0])          # never file contents
        self.assertEqual(found[0].source, "Gemini")
        self.assertGreater(self.service.ai_spent(), 0)
        self.service.decline_suggestion(found[0])
        waiting = self.service.catalog_suggestions()
        self.assertNotIn(found[0].key, {s.key for s in waiting})
        self.assertIn(found[1].key, {s.key for s in waiting})


class SessionFlowTest(unittest.TestCase):
    """An organizing session on copies of the test folders: Choose, Duplicates, Catalog, Review, Move."""

    def setUp(self):
        import shutil
        from tests.fixtures import shared_test_folders

        self.dir = tempfile.TemporaryDirectory()
        base = Path(self.dir.name)
        for name in ("Downloads", "Sorted"):
            shutil.copytree(shared_test_folders() / name, base / name)
        self.base = base
        self.service = AppService(AppPaths(base / "data"), ApiKeyStore(FakeKeyring()))
        self.service.scanner.protected = []
        self.service.add_source(str(base / "Downloads"))
        self.service.add_destination(str(base / "Sorted"))

    def tearDown(self):
        self.dir.cleanup()

    def test_sessions_are_named_saved_listed_and_reopened(self):
        flow = self.service.new_session("Downloads clean-up")
        flow.use_folders()
        other = self.service.new_session("Downloads clean-up")
        self.assertNotEqual(flow.session.file, other.session.file)
        other.rename("Old laptop backup")
        names = [s.name for s, _ in self.service.recent_sessions()]
        self.assertEqual(names, ["Old laptop backup", "Downloads clean-up"])        # the latest used first
        self.assertEqual(self.service.recent_sessions()[0][1], "Step 1 Choose")
        self.service.remove_folder(str(self.base / "Sorted"))
        again = self.service.open_session(flow.session.file)
        self.assertEqual(again.restore_folders(), [])
        self.assertEqual(self.service.destination_folders(), [str(self.base / "Sorted")])   # its folders come back

    def test_label_ideas_merge_and_exact_labels(self):
        ideas = self.service.label_ideas()
        self.assertIn("Work", ideas)
        self.assertNotIn("Documents", ideas)                   # general names are left out
        self.service.add_label("Tax")
        self.service.add_label("Taxes")
        receipt = str(self.base / "Downloads" / "a.txt")
        self.service.set_file_labels([receipt], ["Tax", "Work"])
        self.service.merge_label("Tax", "Taxes")
        self.assertEqual(self.service.labels(), ["Taxes"])
        self.assertEqual(self.service.labels_of(receipt), ["Taxes", "Work"])

    def test_wizard_duplicates_batches_learning_review_and_move(self):
        flow = self.service.new_session("Downloads clean-up")
        for name in self.service.label_ideas()[:8]:
            self.service.add_label(name)
        flow.use_folders()
        flow.set_plan(self.service.make_plan())
        copies = flow.copies()
        self.assertTrue(copies)
        flow.keep_by("oldest")
        oldest = copies[0].kept
        self.assertTrue(all(oldest.modified_ns <= c.modified_ns for c in copies[0].copies))
        flow.keep_by("newest")
        extras = flow.ticked_copies()
        self.assertGreater(extras, 0)
        result = self.service.queue_copies(copies)
        flow.copies_queued(result)
        self.assertEqual(flow.session.copies_queued, result.moved)
        gone = {os.path.normcase(old) for old, _ in result.moves}

        batches = flow.make_batches()
        self.assertTrue(batches)
        paths = [p for b in batches for p in b.paths]
        self.assertEqual(len(paths), len(set(paths)))                         # every file in one batch
        self.assertFalse(gone & {os.path.normcase(p) for p in paths})          # the extra copies have left
        self.assertEqual([b.certainty for b in batches if b.labels],
                         sorted([b.certainty for b in batches if b.labels], reverse=True))   # surest first
        first = flow.current()
        self.assertTrue(first.labels)
        before = flow.files_left()
        unticked = first.paths[-1]
        learned = flow.confirm(first, first.paths[:-1], [first.labels[0][0]])
        self.assertEqual(self.service.labels_of(first.paths[0]), [first.labels[0][0]])
        self.assertLess(flow.files_left(), before)
        self.assertIn(unticked, flow.session.later)                           # unticked: comes back later
        self.assertTrue(flow.is_later(flow.batches[-1]))
        self.assertTrue(flow.back())                                          # Back undoes the answer
        self.assertEqual(flow.current().paths, first.paths)
        self.assertEqual(flow.session.done, [])

        while flow.current() is not None and not flow.learned_enough():       # confirm until SortZen has learned
            b = flow.current()
            if b.labels:
                learned = flow.confirm(b, b.paths, [b.labels[0][0]])
            else:
                flow.send_to_review(b, b.paths)
        self.assertTrue(flow.learned_enough())
        self.assertIsInstance(learned.settled_batches, int)
        b = flow.current()
        if b is not None:
            flow.skip(b)
            self.assertTrue(all(p in flow.session.later for p in b.paths))
        flow.finish_catalog()

        review = flow.start_review(self.service.make_plan())
        self.assertTrue(review)
        self.assertTrue(any(r.sure for r in review))
        self.assertTrue(all(flow.is_reviewed(r) for r in review if r.sure))   # sure batches need no check
        todo = next(r for r in review if not flow.is_reviewed(r) and r.batch.paths)
        s = flow.files_in(todo)[0]
        new = flow.new_folder(str(self.base / "Sorted"), "Jordan")
        self.assertIn(new, flow.planned_folders())
        undo = flow.move_file([s.path], new)
        self.assertEqual(self.service.corrections()[s.path], new)
        self.assertTrue(s.new_folder)
        deleted = flow.delete_folder(new)                                     # its files go to the folder it was in
        self.assertEqual(s.destination, str(self.base / "Sorted"))
        flow.undo_delete_folder(deleted)
        self.assertEqual(s.destination, new)
        flow.undo_move_file(undo)
        with self.assertRaises(ValueError):
            flow.delete_folder(str(self.base / "Sorted"))                      # folders that exist: in Explorer
        for r in review:
            if not flow.is_reviewed(r):
                last = flow.confirm_review(r)
        self.assertTrue(last)
        rows = flow.move_rows()
        self.assertTrue(rows)
        result = self.service.move(flow.plan, rows)
        flow.moved(result)
        self.assertGreater(result.moved, 0)
        self.assertEqual(self.service.recent_sessions()[0][1], f"Moved · {result.moved:,} files")
